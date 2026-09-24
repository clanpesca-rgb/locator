#!/usr/bin/env python3
"""04: Excel del confronto v2 (famiglie, fascia da-a). Richiede openpyxl.
Fogli: Sintesi / Confronto famiglie (una riga per padre) / Dettaglio (tutte le
rilevazioni con link) / Marche. La disponibilita' e' registrata, mai filtrata.
Dopo la scrittura eseguire il ricalcolo (recalc LibreOffice) prima di consegnare."""
import csv, os, statistics, sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = "24/09/2026"
OUT = os.path.join(REPO, f"confronto_prezzi_2026-09-24.xlsx")

F = "Arial"
TITLE = Font(name=F, size=15, bold=True, color="1F3864")
HDR_FONT = Font(name=F, size=10, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", start_color="1F3864")
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
LINK = Font(name=F, size=10, color="0563C1", underline="single")
NOTE = Font(name=F, size=9, italic=True, color="595959")
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
EUR = "€ #,##0.00"
PCT = "+0.0%;-0.0%;0.0%"

def header_row(ws, row, headers, widths):
    for c, (h, wd) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border = HDR_FONT, HDR_FILL, THIN
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = wd
    ws.row_dimensions[row].height = 26

rows = list(csv.DictReader(open(os.path.join(REPO, "dati", "confronto_v2.csv"))))
catalogo = list(csv.DictReader(open(os.path.join(REPO, "dati", "catalogo_padri.csv"))))
tot_marca = defaultdict(int)
for p in catalogo:
    tot_marca[p["marca"]] += 1

# per padre: miglior (minimo) prezzo web
best = {}
for r in rows:
    pv = float(r["prezzo_concorrente"])
    k = r["id"]
    if k not in best or pv < float(best[k]["prezzo_concorrente"]):
        best[k] = r

wb = Workbook()

# ---------------- Dettaglio
wsd = wb.create_sheet("Dettaglio")
header_row(wsd, 1, ["ID", "Marca", "Prodotto (padre)", "Categoria", "CP da", "CP a",
                    "Sito", "Prezzo web", "Fascia web da", "Fascia web a", "Disponibilità",
                    "Livello match", "Delta vs fascia", "Link concorrente", "Link Clan Pesca"],
           [8, 13, 42, 16, 10, 10, 24, 11, 10, 10, 12, 11, 11, 16, 14])
r = 2
for row in sorted(rows, key=lambda x: (x["marca"], x["prodotto"], x["sito"])):
    vals = [row["id"], row["marca"], row["prodotto"], row["categoria"],
            float(row["cp_prezzo_da"]), float(row["cp_prezzo_a"]), row["sito"],
            float(row["prezzo_concorrente"]),
            float(row["conc_da"]) if row["conc_da"] else None,
            float(row["conc_a"]) if row["conc_a"] else None,
            row["disponibilita_concorrente"] or "n.d.", row["livello_match"]]
    for c, v in enumerate(vals, 1):
        cell = wsd.cell(row=r, column=c, value=v)
        cell.font, cell.border = BASE, THIN
    for c in (5, 6, 8, 9, 10):
        wsd.cell(row=r, column=c).number_format = EUR
    d = wsd.cell(row=r, column=13, value=f"=IF(H{r}<E{r},(E{r}-H{r})/H{r},IF(H{r}>F{r},(F{r}-H{r})/H{r},0))")
    d.number_format, d.font, d.border = PCT, BASE, THIN
    for c, (txt, url) in ((14, ("apri", row["url_concorrente"])), (15, ("apri", row["url_clanpesca"]))):
        u = wsd.cell(row=r, column=c, value=txt)
        u.hyperlink, u.font, u.border = url, LINK, THIN
    r += 1
last_d = r - 1
wsd.freeze_panes = "A2"
wsd.auto_filter.ref = f"A1:O{last_d}"
wsd.conditional_formatting.add(f"M2:M{last_d}", CellIsRule(operator="greaterThan", formula=["0.02"], font=Font(name=F, size=10, color="C00000")))
wsd.conditional_formatting.add(f"M2:M{last_d}", CellIsRule(operator="lessThan", formula=["-0.02"], font=Font(name=F, size=10, color="006100")))

# ---------------- Confronto famiglie (una riga per padre)
wsf = wb.create_sheet("Confronto famiglie")
header_row(wsf, 1, ["ID", "Marca", "Prodotto (padre)", "Clan Pesca da", "Clan Pesca a",
                    "N. siti con prezzo", "Miglior prezzo web", "Sito del migliore",
                    "Posizione Clan Pesca", "Delta vs fascia"],
           [8, 13, 46, 12, 12, 10, 12, 24, 16, 11])
r = 2
n_sites = defaultdict(set)
for row in rows:
    n_sites[row["id"]].add(row["sito"])
for pid, row in sorted(best.items(), key=lambda kv: (kv[1]["marca"], kv[1]["prodotto"])):
    cp_da, cp_a = float(row["cp_prezzo_da"]), float(row["cp_prezzo_a"])
    vals = [pid, row["marca"], row["prodotto"], cp_da, cp_a, len(n_sites[pid]),
            float(row["prezzo_concorrente"]), row["sito"]]
    for c, v in enumerate(vals, 1):
        cell = wsf.cell(row=r, column=c, value=v)
        cell.font, cell.border = BASE, THIN
    for c in (4, 5, 7):
        wsf.cell(row=r, column=c).number_format = EUR
    pos = wsf.cell(row=r, column=9, value=f'=IF(G{r}<D{r},"più caro",IF(G{r}>E{r},"più economico","in fascia"))')
    pos.font, pos.border = BASE, THIN
    d = wsf.cell(row=r, column=10, value=f"=IF(G{r}<D{r},(D{r}-G{r})/G{r},IF(G{r}>E{r},(E{r}-G{r})/G{r},0))")
    d.number_format, d.font, d.border = PCT, BASE, THIN
    r += 1
last_f = r - 1
wsf.freeze_panes = "A2"
wsf.auto_filter.ref = f"A1:J{last_f}"
wsf.conditional_formatting.add(f"J2:J{last_f}", CellIsRule(operator="greaterThan", formula=["0.02"], font=Font(name=F, size=10, color="C00000")))
wsf.conditional_formatting.add(f"J2:J{last_f}", CellIsRule(operator="lessThan", formula=["-0.02"], font=Font(name=F, size=10, color="006100")))

# ---------------- Marche
wsm = wb.create_sheet("Marche")
header_row(wsm, 1, ["Marca", "Padri a catalogo", "Padri con prezzo web", "Delta mediano vs fascia",
                    "Più caro", "In fascia", "Più economico"], [16, 12, 12, 13, 9, 9, 12])
per_brand = defaultdict(list)
for pid, row in best.items():
    cp_da, cp_a = float(row["cp_prezzo_da"]), float(row["cp_prezzo_a"])
    pv = float(row["prezzo_concorrente"])
    delta = (cp_da - pv) / pv * 100 if pv < cp_da else ((cp_a - pv) / pv * 100 if pv > cp_a else 0.0)
    per_brand[row["marca"]].append(delta)
r = 2
for b in sorted(tot_marca, key=lambda x: -tot_marca[x]):
    ds = per_brand.get(b, [])
    vals = [b, tot_marca[b], len(ds),
            statistics.median(ds) / 100 if ds else None,
            sum(1 for d in ds if d > 2), sum(1 for d in ds if -2 <= d <= 2), sum(1 for d in ds if d < -2)]
    for c, v in enumerate(vals, 1):
        cell = wsm.cell(row=r, column=c, value=v)
        cell.font, cell.border = (BOLD if c == 1 else BASE), THIN
    wsm.cell(row=r, column=4).number_format = PCT
    r += 1
wsm.cell(row=r + 1, column=1, value="Delta calcolato sul MIGLIOR prezzo web di ogni prodotto rispetto alla fascia Clan Pesca da-a; "
        "valori del 24/09/2026 (fotografia). Righe di dettaglio e formule nel foglio Dettaglio.").font = NOTE
wsm.freeze_panes = "A2"

# ---------------- Sintesi
wss = wb.active
wss.title = "Sintesi"
wss.column_dimensions["A"].width = 3
wss.column_dimensions["B"].width = 56
wss.column_dimensions["C"].width = 15
wss["B2"] = "Confronto prezzi Clan Pesca vs web — famiglie di prodotto"
wss["B2"].font = TITLE
wss["B3"] = f"Rilevazione del {DATA} · unità = prodotto padre con fascia da-a · disponibilità ignorata (conta solo il prezzo) · esclusi Shimano e Tuning"
wss["B3"].font = NOTE
tot_conf = len(best)
piu_caro = sum(1 for pid, row in best.items()
               if float(row["prezzo_concorrente"]) < float(row["cp_prezzo_da"]))
in_fascia = sum(1 for pid, row in best.items()
                if float(row["cp_prezzo_da"]) <= float(row["prezzo_concorrente"]) <= float(row["cp_prezzo_a"]))
piu_econ = tot_conf - piu_caro - in_fascia
kv = [
    ("Prodotti padre a catalogo (senza Shimano/Tuning)", len(catalogo)),
    ("— con almeno un prezzo trovato sul web", tot_conf),
    ("Righe di confronto totali", None),
    ("Clan Pesca PIÙ CARA del miglior prezzo web", piu_caro),
    ("Fascia Clan Pesca comprende il miglior prezzo web", in_fascia),
    ("Clan Pesca PIÙ ECONOMICA del miglior prezzo web", piu_econ),
]
r = 5
for label, val in kv:
    wss.cell(row=r, column=2, value=label).font = BASE
    c = wss.cell(row=r, column=3)
    c.value = f"=COUNTA(Dettaglio!A2:A{last_d})" if val is None else val
    c.font = BOLD
    r += 1
r += 1
for line in [
    "Come leggere: 'più caro' = il miglior prezzo web è sotto la fascia Clan Pesca; 'in fascia' = cade dentro la fascia da-a;",
    "'più economico' = il web costa più del massimo della fascia. Delta calcolati rispetto al bordo di fascia più vicino.",
    "Livelli di match: esatto > famiglia > probabile > possibile (colonna nel Dettaglio; i 'possibile' vanno verificati con un click).",
    "Siti non verificabili in automatico (bot-wall): pescaloccasione, amazon, trovaprezzi, idealo, decathlon, marcosportshop, free-fishing.",
]:
    wss.cell(row=r, column=2, value=line).font = NOTE
    r += 1

wb.save(OUT)
print("scritto", OUT, f"({last_d-1} righe dettaglio, {tot_conf} famiglie)")

# copia i CSV finali dall'area di lavoro (ignorata da git) alla cartella versionata
import shutil
shutil.copy(os.path.join(REPO, "dati", "confronto_v2.csv"), os.path.join(REPO, "confronto_famiglie.csv"))
shutil.copy(os.path.join(REPO, "dati", "catalogo_padri.csv"), os.path.join(REPO, "catalogo_padri.csv"))
print("copiati confronto_famiglie.csv e catalogo_padri.csv in ricerca_prezzi/")
